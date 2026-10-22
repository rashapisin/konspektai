"""Разбор расписания и плана-конспекта на искусственных документах того же вида, что и настоящие."""

import docx
import pytest

from stv.ai import build_prompt, clean_answer
from stv.documents import match_plan, parse_plan, parse_schedule
from stv.engine.journal import Entry


def _schedule(path, special=False):
    d = docx.Document()
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "УТВЕРЖДАЮ\nНачальник отряда\n«___» ____ 2026 г."
    d.add_paragraph("Расписание занятий по специальной подготовке на апрель 2026 года" if special
                    else "Расписание занятий на апрель 2026 года")
    d.add_paragraph("с общей группой")
    t = d.add_table(rows=4, cols=4)
    for i, h in enumerate(("Дата", "Часы занятий", "Предмет обучения, тема и учебные вопросы занятий", "Метод")):
        t.rows[0].cells[i].text = h
    for i in range(4):
        t.rows[1].cells[i].text = str(i + 1)
    t.rows[2].cells[0].text = "08.04.2026"
    t.rows[2].cells[2].text = ("Психологическая подготовка\nТема № 2. Последствия хронического стресса.\n"
                               "Эмоциональное выгорание.\nУчебный вопрос № 1 Последствия (35 мин).\n"
                               "Учебный вопрос №2 – Выгорание (35 мин).")
    t.rows[3].cells[0].text = "09.04.2026"
    t.rows[3].cells[2].text = "Тема № 10 Требование к оборудованию мест проведения соревнований"
    d.add_paragraph("Инженер отделения")
    d.save(path)


def _plan(path):
    d = docx.Document()
    for t in ("УТВЕРЖДАЮ", "Начальник отряда", "ПЛАН-КОНСПЕКТ", "по служебной подготовке",
              "Тема: Последствия хронического стресса. Эмоциональное выгорание.",
              "Цель занятия: Рассмотрение последствий стресса.", "Литература:", "1. Бойко В.В. Синдром выгорания.",
              "Развернутый план занятия:", "Учебный вопрос №1 – Последствия хронического стресса (35 мин).",
              "Хронический стресс действует разрушительно."):
        d.add_paragraph(t)
    d.add_paragraph("снижение продуктивности в работе", style="List Bullet")
    d.add_paragraph("2. Учебный вопрос №2 – Эмоциональное выгорание (35 мин).")
    d.add_paragraph("Эмоциональное выгорание - комплекс последствий хронического стресса.")
    d.add_paragraph("Пособия и оборудование, используемые на занятии: ______")
    d.add_paragraph("Инженер отделения ОСПиП")
    d.save(path)


def test_service_schedule(tmp_path):
    f = tmp_path / "s.docx"
    _schedule(f)
    les = parse_schedule(f)
    assert [l.date for l in les] == ["08.04.26", "09.04.26"]
    assert les[0].kind == "service"
    assert les[0].subject == "Психологическая подготовка"
    assert les[0].topic == "Последствия хронического стресса. Эмоциональное выгорание."
    assert len(les[0].questions) == 2
    assert les[1].topic == "Требование к оборудованию мест проведения соревнований."


def test_special_schedule_kind(tmp_path):
    f = tmp_path / "s.docx"
    _schedule(f, special=True)
    assert {l.kind for l in parse_schedule(f)} == {"special"}


def test_plan_is_cleaned(tmp_path):
    f = tmp_path / "p.docx"
    _plan(f)
    p = parse_plan(f)
    assert p.topic == "Последствия хронического стресса. Эмоциональное выгорание."
    assert p.goal.startswith("Рассмотрение")
    assert p.text.startswith("Учебный вопрос №1")
    assert "- снижение продуктивности в работе" in p.text
    assert "Учебный вопрос №2" in p.text
    for junk in ("УТВЕРЖДАЮ", "Бойко", "Пособия", "Инженер"):
        assert junk not in p.text


def test_plan_matches_lesson(tmp_path):
    f = tmp_path / "p.docx"
    _plan(f)
    p = parse_plan(f)
    assert match_plan("Последствия хронического стресса. Эмоциональное выгорание.", [p]) is p
    assert match_plan("Требование к оборудованию мест проведения соревнований.", [p]) is None


def test_prompt_and_answer_cleanup():
    prompt = build_prompt(Entry(kind="service", topic="Выгорание"), "Исходный текст.", 1400)
    assert "1400" in prompt and "Исходный текст." in prompt and "Выгорание" in prompt
    answer = "Вот конспект:\n**Учебный вопрос №1**\n* пункт один\n# лишний заголовок"
    assert clean_answer(answer) == "Учебный вопрос №1\n- пункт один\nлишний заголовок"


def test_old_word_doc_text():
    """Старый Word (.doc): текст собирается по «таблице кусков» — UTF-16 и 8-битные куски, служебные символы."""
    import struct
    from stv.documents import doc_stream_text
    part1 = "Тема: Огнетушители\rВиды: \x13 FORMTEXT \x14порошковые\x15, углекислотные\r".encode("utf-16-le")
    part2 = b"OK\x07"
    fib = bytearray(1024)
    struct.pack_into("<H", fib, 0, 0xA5EC)
    struct.pack_into("<H", fib, 32, 14)          # csw
    struct.pack_into("<H", fib, 62, 22)          # cslw
    blob = 32 + 2 + 28 + 2 + 88 + 2
    text_at = 600
    fib[text_at:text_at + len(part1)] = part1
    fib[900:900 + len(part2)] = part2
    n1 = len(part1) // 2
    plc = struct.pack("<3I", 0, n1, n1 + len(part2))
    plc += struct.pack("<HIH", 0, text_at, 0) + struct.pack("<HIH", 0, (900 * 2) | 0x40000000, 0)
    clx = b"\x02" + struct.pack("<I", len(plc)) + plc
    table = b"\x00" * 16 + clx
    struct.pack_into("<II", fib, blob + 33 * 8, 16, len(clx))
    text = doc_stream_text(bytes(fib), table)
    assert text == "Тема: Огнетушители\nВиды: порошковые, углекислотные\nOK"
